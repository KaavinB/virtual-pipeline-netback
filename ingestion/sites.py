"""Choose example sites from real satellite-detected flares and load core.dim_site.

    python -m ingestion.sites             # download (cached), select, upsert
    python -m ingestion.sites --dry-run   # print the selection without writing

Source: World Bank / NOAA VIIRS "Flare Volume Estimates by Individual Flare Location
2012-2025". Basin membership uses EIA's published boundaries (data/reference/*.geojson).
The data says a flare exists; it does not say the site lacks a pipeline. "Off-pipeline" is an
assumption, made more plausible by keeping only flares that burned gas in each of the last
three years: a one-off upset flares once, a site with no takeaway flares every year.
"""

import argparse
import json
import logging
import math
import sys
from dataclasses import dataclass
from pathlib import Path

import openpyxl
import requests

from ingestion.config import ROOT, setup_logging
from ingestion.db import connect

log = logging.getLogger("sites")

WORKBOOK_URL = (
    "https://thedocs.worldbank.org/en/doc/b34e0c054bb3fe3695e70154c28eef3f-0400072026/related/"
    "Flare-Volume-Estimates-by-individual-Flare-Location-2012-2025.xlsx"
)
WORKBOOK_PATH = ROOT / "data" / "raw" / "flare_volumes_2012_2025.xlsx"
WORKBOOK_COLUMNS = ("id", "lat", "lon", "Country", "MMSCFD", "Year", "Field type", "Field name")
BASINS = {  # basin -> (EIA boundary file, number of sites, site-name prefix)
    "Permian": (ROOT / "data" / "reference" / "permian_basin_eia.geojson", 12, "PB"),
    "Eagle Ford": (ROOT / "data" / "reference" / "eagle_ford_play_eia.geojson", 6, "EF"),
}

# ── selection rule (see docs/05-sites-and-routes.md) ─────────────────────────
LATEST_YEAR = 2025
PERSISTENCE_YEARS = (2023, 2024, 2025)  # must have flared in every one of these
FIELD_TYPE = "OIL"  # associated gas: the stranded-gas case
MIN_MMSCFD = 0.10  # below this, a trailer would take a week or more to fill
MAX_MMSCFD = 2.00  # above this, a gathering line usually pencils; not a trucking customer
MIN_SPACING_KM = 20  # spread sites out so routes cover a range of distances

# Texas / New Mexico: the border follows 32°N westward from 103.064°W (Wikipedia, approx).
NM_BORDER_LON, NM_BORDER_LAT = -103.064, 32.0


@dataclass
class Flare:
    flare_id: str
    lat: float
    lon: float
    field_type: str | None
    field_name: str | None
    mmscfd: dict[int, float]


@dataclass(frozen=True)
class Site:
    flare_id: str
    site_name: str
    basin: str
    state: str
    lat: float
    lon: float
    field_type: str | None
    field_name: str | None
    flared_mmscfd: float
    years_flared_of_3: int
    selection_rank: int


def download_workbook(path: Path = WORKBOOK_PATH) -> Path:
    if path.exists():
        return path
    path.parent.mkdir(parents=True, exist_ok=True)
    log.info("downloading World Bank flare workbook (~47 MB)")
    with requests.get(WORKBOOK_URL, stream=True, timeout=120) as resp:
        resp.raise_for_status()
        tmp = path.with_suffix(".part")
        with tmp.open("wb") as fh:
            for chunk in resp.iter_content(1 << 20):
                fh.write(chunk)
    tmp.rename(path)  # only a complete download gets the real name
    return path


def read_us_flares(path: Path, years=PERSISTENCE_YEARS) -> dict[str, Flare]:
    """One Flare per id, with MMSCFD for each year of interest (US rows only)."""
    wb = openpyxl.load_workbook(path, read_only=True)
    rows = wb["Dashboard_file"].iter_rows(values_only=True)
    header = next(rows)
    col = {name: header.index(name) for name in WORKBOOK_COLUMNS}
    flares: dict[str, Flare] = {}
    for r in rows:
        if r[col["Country"]] != "United States" or r[col["Year"]] not in years:
            continue
        f = flares.setdefault(
            r[col["id"]],
            Flare(r[col["id"]], float(r[col["lat"]]), float(r[col["lon"]]),
                  r[col["Field type"]], r[col["Field name"]], {}),
        )  # fmt: skip
        f.mmscfd[int(r[col["Year"]])] = float(r[col["MMSCFD"]] or 0)
    wb.close()
    return flares


def load_ring(path: Path) -> list[tuple[float, float]]:
    """Outer ring of a single-polygon GeoJSON as (lon, lat) pairs."""
    geometry = json.loads(path.read_text())["features"][0]["geometry"]
    assert geometry["type"] == "Polygon", geometry["type"]
    return [(lon, lat) for lon, lat in geometry["coordinates"][0]]


def point_in_polygon(lon: float, lat: float, ring: list[tuple[float, float]]) -> bool:
    """Ray casting: cast a ray east from the point; an odd number of edge crossings = inside."""
    inside = False
    for (x1, y1), (x2, y2) in zip(ring, ring[1:] + ring[:1], strict=True):
        if (y1 > lat) != (y2 > lat):  # edge straddles the ray's latitude
            x_cross = x1 + (lat - y1) * (x2 - x1) / (y2 - y1)
            if x_cross > lon:
                inside = not inside
    return inside


def haversine_km(lat1: float, lon1: float, lat2: float, lon2: float) -> float:
    p1, p2 = math.radians(lat1), math.radians(lat2)
    a = (math.sin((p2 - p1) / 2) ** 2
         + math.cos(p1) * math.cos(p2) * math.sin(math.radians(lon2 - lon1) / 2) ** 2)  # fmt: skip
    return 2 * 6371.0 * math.asin(math.sqrt(a))


def state_of(lat: float, lon: float) -> str:
    return "NM" if lat > NM_BORDER_LAT and lon < NM_BORDER_LON else "TX"


def select(flares: dict[str, Flare], basin: str, ring, count: int, prefix: str) -> list[Site]:
    candidates = [
        f for f in flares.values()
        if f.field_type == FIELD_TYPE
        and all(f.mmscfd.get(y, 0) > 0 for y in PERSISTENCE_YEARS)
        and MIN_MMSCFD <= f.mmscfd[LATEST_YEAR] <= MAX_MMSCFD
        and point_in_polygon(f.lon, f.lat, ring)
    ]  # fmt: skip
    # Greedy spread: largest flares first, skipping any within MIN_SPACING_KM of a pick.
    # Ties are broken by id so the selection is deterministic.
    candidates.sort(key=lambda f: (-f.mmscfd[LATEST_YEAR], f.flare_id))
    picked: list[Flare] = []
    for f in candidates:
        if all(haversine_km(f.lat, f.lon, p.lat, p.lon) >= MIN_SPACING_KM for p in picked):
            picked.append(f)
        if len(picked) == count:
            break
    log.info("%s: %d candidates pass the rule, %d picked", basin, len(candidates), len(picked))
    return [
        Site(f.flare_id, f"{prefix}-{i:02d}", basin, state_of(f.lat, f.lon),
             round(f.lat, 5), round(f.lon, 5), f.field_type, f.field_name,
             round(f.mmscfd[LATEST_YEAR], 4), len(PERSISTENCE_YEARS), i)
        for i, f in enumerate(picked, start=1)
    ]  # fmt: skip


def save(sites: list[Site]) -> None:
    with connect() as conn:
        conn.execute(
            "DELETE FROM core.dim_site WHERE flare_id <> ALL(%s)", ([s.flare_id for s in sites],)
        )  # dim_site mirrors the current selection; routes cascade
        conn.cursor().executemany(
            """INSERT INTO core.dim_site
                   (flare_id, site_name, basin, state, lat, lon, field_type, field_name,
                    flared_mmscfd, flare_year, years_flared_of_3, selection_rank)
               VALUES (%(flare_id)s, %(site_name)s, %(basin)s, %(state)s, %(lat)s, %(lon)s,
                       %(field_type)s, %(field_name)s, %(flared_mmscfd)s, %(flare_year)s,
                       %(years_flared_of_3)s, %(selection_rank)s)
               ON CONFLICT (flare_id) DO UPDATE SET
                   site_name = EXCLUDED.site_name, basin = EXCLUDED.basin,
                   state = EXCLUDED.state, lat = EXCLUDED.lat, lon = EXCLUDED.lon,
                   field_type = EXCLUDED.field_type, field_name = EXCLUDED.field_name,
                   flared_mmscfd = EXCLUDED.flared_mmscfd, flare_year = EXCLUDED.flare_year,
                   years_flared_of_3 = EXCLUDED.years_flared_of_3,
                   selection_rank = EXCLUDED.selection_rank, loaded_at = now()""",
            [{**s.__dict__, "flare_year": LATEST_YEAR} for s in sites],
        )
        conn.execute("ANALYZE core.dim_site")  # keep planner estimates current
    log.info("dim_site: %d sites saved", len(sites))


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args(argv)
    setup_logging()

    flares = read_us_flares(download_workbook())
    log.info("read %d US flares with data in %s", len(flares), PERSISTENCE_YEARS)
    sites = [
        site
        for basin, (boundary, count, prefix) in BASINS.items()
        for site in select(flares, basin, load_ring(boundary), count, prefix)
    ]
    for s in sites:
        log.info("%s %s %-5.2f MMSCFD  (%.4f, %.4f)  %s", s.site_name, s.state,
                 s.flared_mmscfd, s.lat, s.lon, s.field_name or "")  # fmt: skip
    if not args.dry_run:
        save(sites)
    return 0


if __name__ == "__main__":
    sys.exit(main())
